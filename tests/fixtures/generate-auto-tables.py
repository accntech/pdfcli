"""Synthetic automatic-detection regression fixtures; no user data."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table
from openpyxl.drawing.image import Image
from PIL import Image as Pixels
from io import BytesIO

book = Workbook()
book.remove(book.active)
for name, right, left_start, right_start in [
    ("Gutter", 6, 1, 1),
    ("Adjacent", 5, 5, 2),
    ("Shared heading", 6, 2, 2),
]:
    sheet = book.create_sheet(name)
    for column in range(1, right + 4):
        sheet.column_dimensions[get_column_letter(column)].width = 22
    if name == "Adjacent":
        sheet.merge_cells("A4:D4")
        sheet["A4"] = "LEFT TABLE HEADING"
        sheet.merge_cells("E1:H1")
        sheet["E1"] = "RIGHT TABLE HEADING"
    elif name == "Shared heading":
        sheet.merge_cells("A1:I1")
        sheet["A1"] = "SHARED TABLE HEADING"
    for side, col, start in [("LEFT", 1, left_start), ("RIGHT", right, right_start)]:
        for number in range(1, 91):
            row = number + start - 1
            sheet.cell(row, col, f"{side} ROW {number:02}")
            sheet.cell(row, col + 3, f"{side} VALUE {number:02}")
            sheet.row_dimensions[row].height = 24
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.scale = 100
    sheet.print_area = "A1:B10"  # Deliberately stale: default must retain all data.
book.save(Path(__file__).with_name("excel-auto-tables.xlsx"))

objects = Workbook()
sheet = objects.active
sheet.title = "Structured tables"
for col in range(1, 9):
    sheet.column_dimensions[get_column_letter(col)].width = 22
    sheet.cell(1, col, f"HEADER{col}")
for row in range(2, 92):
    sheet.row_dimensions[row].height = 24
    for side, column in [("LEFT", 1), ("RIGHT", 5)]:
        sheet.cell(row, column, f"{side} ROW {row - 1:02}")
        sheet.cell(row, column + 3, f"{side} VALUE {row - 1:02}")
sheet.add_table(Table(displayName="LeftTable", ref="A1:D91"))
sheet.add_table(Table(displayName="RightTable", ref="E1:H91"))
sheet = objects.create_sheet("Spanning image")
for col in range(1, 10):
    sheet.column_dimensions[get_column_letter(col)].width = 15
for cell, value in [("A25", "GRAPHIC LEFT"), ("D25", "LEFT VALUE"),
                    ("F25", "GRAPHIC RIGHT"), ("I25", "RIGHT VALUE")]:
    sheet[cell] = value
png = BytesIO()
Pixels.new("RGB", (1200, 120), "#497eaf").save(png, format="PNG")
png.seek(0)
sheet.add_image(Image(png), "A1")
for sheet in objects:
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.scale = 100
objects.save(Path(__file__).with_name("excel-auto-objects.xlsx"))
